"""Self-contained current workbooks and reports; no 486-source runtime dependency."""
from common import *
import pandas as pd,numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from potential_delivery import POTENTIAL_CN,main_text,experiment_text
CN={'category_id':'主题ID','name':'主题名称','report_rank':'榜内排序','numeric_rank':'数值候选排序','family':'热点类型','scenario':'情景','kind':'情景类型','core_score':'多年核心评分','emerging_score':'多年新兴评分','core_papers':'核心窗口论文数','core_years':'核心窗口年数','recent_papers':'最近一年论文数','previous_papers':'前一年论文数','baseline_papers':'此前多年论文数','baseline_annual_mean_papers':'此前多年年均论文数','baseline_years_used':'基线年数','multiyear_share_ratio':'最近一年对多年均值份额比','share_growth_ratio':'最近一年同比份额比','historical_peak_ratio':'最近一年对前四年最高份额比','core_active_quarters':'核心窗口活跃季度数','core_active_years':'核心窗口活跃年份数','core_institutions':'核心窗口去重机构数','recent_institutions':'最近一年去重机构数','institution_expansion':'最近一年机构数对基线年均之比','three_year_log_share_slope':'近三年对数份额年斜率','growing_quarters':'最近一年同比增长季度数','robust_min_ratio':'数据确认最小份额比','share_ratio_2026_ytd':'2026年1至8月同期份额比','core_review_pass':'核心样本范围审阅通过','emerging_review_pass':'新兴样本范围审阅通过','final_core':'核心初评入选','final_emerging':'新兴初评入选','core_eligible':'核心数值门槛通过','emerging_robust':'新兴数值及数据确认通过','paper_review_reason':'样本审阅依据与限制','review_status':'评审状态','trajectory':'五年轨迹解释','parameters':'参数','numeric_candidates':'数值候选数','reviewed_selected':'固定审阅下入选数','baseline_count':'本版基准入选数','retained':'本版保留数','retention':'本版保留比例','jaccard':'集合Jaccard','gained':'增加ID','lost':'退出ID','selected_ids':'情景入选ID','interpretation':'解释限制','spearman':'排名Spearman相关','weights':'权重','pool_size':'排名池大小','mean_abs_rank_change':'平均名次变动','max_abs_rank_change':'最大名次变动','old_selected':'修订前入选','new_selected':'修订后入选','change':'名单变化','old_score':'旧定义评分','new_score':'新定义评分','note':'说明','reason':'范围判断依据','decision':'范围判断','row_id':'文献行ID','title':'题名','excerpt':'实际展示摘录','annual_block':'年度区间','review_purpose':'补审目的','data_scenarios_passed':'五项数据压力测试通过数','robustness_grade':'数据压力测试结论','core_review_basis':'核心审阅来源','publication_year':'年份','papers':'论文数','observed_months':'覆盖月份数','title_only_fraction':'仅题名比例','jan1_proxy_fraction':'1月1日日期比例','metadata_match_fraction':'元数据关联比例','institution_metadata_fraction':'机构信息非空比例','start_quarter':'年度起始季度','end_quarter':'年度截止季度','share':'年度平滑份额','background':'年度论文背景总数','block':'距最近年度块数','status':'类别状态','analysis_scope':'领域范围','domain':'领域','score':'情景评分','end_quarter':'截止季度','core_mean_annual_share':'核心窗口年均份额','baseline_mean_share':'此前多年年均份额','citation_cohort_percentile':'同发表年引用百分位均值','scope':'解释范围','read_records':'本次阅读条数','read_row_ids':'本次已读文献行ID','reviewer':'审阅方式'}
CN['time_sensitivity_note']='时间窗口敏感性提示'
CN.update(POTENTIAL_CN)
for i in range(5):
 for key,label in [('papers','论文数'),('denominator','背景论文数'),('share','份额')]:CN[f'year{i}_{key}']=f'距最近{i}年度_{label}'
for s,label in [('quality','文本日期质量'),('geometry','原几何筛查'),('dedup','题名年度去重'),('exclude_needs_review','排除待复核'),('supported_only','仅规则支持')]:CN[s+'_retained']=label+'是否保留'

def safe(v):
 if isinstance(v,(dict,list,tuple)):v=json.dumps(v,ensure_ascii=False)
 if v is None or (not isinstance(v,str) and pd.isna(v)):return None
 if isinstance(v,np.generic):v=v.item()
 if isinstance(v,str):
  v=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',v)[:30000]
  if v[:1] in '=+-@':v="'"+v
 return v

def table(df,cols):
 rows=['| '+' | '.join(CN.get(c,c) for c in cols)+' |','| '+' | '.join(['---']*len(cols))+' |']
 for row in df[cols].itertuples(index=False,name=None):rows.append('| '+' | '.join(('—' if pd.isna(v) else f'{v:.3f}' if isinstance(v,float) else str(v)).replace('|','/').replace('\n',' ') for v in row)+' |')
 return '\n'.join(rows)

def deliver():
 R=BASE/'results';E=BASE/'reliability';z=pd.read_csv(R/'hotspot_summary_all750.csv');core=pd.read_csv(R/'core_hotspots.csv');em=pd.read_csv(R/'emerging_hotspots.csv');pot=pd.read_csv(R/'potential_priority.csv');g=pd.read_csv(E/'gate_and_data_sensitivity.csv');wt=pd.read_csv(E/'weight_trials.csv');abl=pd.read_csv(E/'score_ablation_summary.csv');cov=pd.read_csv(R/'multiyear_coverage_audit.csv');changes=pd.read_csv(R/'multiyear_before_after.csv')
 manifest=[]
 def make(file,notes,tables):
  wb=Workbook();wb.remove(wb.active);spec=[]
  def sheet(name,df):
   ws=wb.create_sheet(name);ws.append([CN.get(c,c) for c in df]);ws.freeze_panes='C2'
   for row in df.itertuples(index=False,name=None):ws.append([safe(x) for x in row])
   ws.auto_filter.ref=ws.dimensions
   for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='164A63');cell.alignment=Alignment(wrap_text=True)
   for i,col in enumerate(df,1):ws.column_dimensions[get_column_letter(i)].width=65 if any(x in col for x in ['reason','excerpt','interpretation']) else 32 if col in ['name','scenario','title','review_status','trajectory'] else 20
   for row in ws.iter_rows(min_row=2):
    for cell in row:
     if isinstance(cell.value,float):cell.number_format='0.0000'
  sheet('阅读说明',pd.DataFrame(notes,columns=['项目','说明']))
  for name,path,cols in tables:
   df=pd.read_csv(BASE/path)
   if cols:df=df[cols]
   sheet(name,df);spec.append(dict(sheet=name,file=path,columns=list(df.columns)))
  sheet('字段对照',pd.DataFrame(list(CN.items()),columns=['字段','中文释义']))
  wb.save(BASE/file);manifest.append(dict(workbook=file,tables=spec))
 notes=[('版本','750-multiyear-v1；固定750标签，核心与新兴已改为多年方法。'),('结果',f'{len(core)}个核心初评候选、{len(em)}个新兴/持续升温初评候选；潜在4任务单元中2项双口径支持、1项条件跟踪、1项观察。'),('核心时间尺度','主窗2023-07至2026-06，共3年；同时要求最近一年活跃。1/3/5年另做敏感性。'),('新兴时间尺度','最近2025-07至2026-06，对比此前2022-07至2025-06三个不重叠年度的平均份额；五年历史2021-07至2026-06。'),('份额与覆盖','各年度分别按同期论文背景归一后平均，避免某个收录量大的年度支配基线；2026年只有1至8月，不与全年直接比总量。'),('核心门槛','三年论文≥750、每年≥150、12季度至少9季度各≥25；最近一年≥250且同比份额比≥0.8；领域与语义审阅通过。'),('新兴门槛','最近一年≥100、前三年年均≥50、多年份额比≥1.25、同比≥1.15、至少3季度同比增长、三年斜率>0、超过前四年峰值≥5%；计数检验和数据方向确认通过。'),('新兴含义','文献关注度持续上升，可能包括成熟技术的新一轮增长，不等于技术首次出现。低基数新苗头及历史峰值以下恢复保留观察。'),('审阅范围','全部750类均计算；核心审阅范围是新评分前25加既有通过项，新兴覆盖所有本次数值稳健候选；其余未审核心明确列为候选待审，不是无价值。'),('样本限制','跨年补审是模型辅助阅读，不是独立专家金标准；方向通过也不等于所有记录纯净。'),('统计限制','计数近似检验未校正所有时间相关与收录偏差；历史截止点使用冻结标签及本次引用快照，属于回顾性敏感性，不是预测回测。'),('潜在范围','保留2026年同集合专利、论文和任务政策口径；没有多年专利数据，不构造专利增长。'),('并列排名','排名时将0—100分取小数点后10位以消除浮点尾差；并列按ID升序展示，实验采用并列最小名次，Spearman也使用舍入后的分数。原始评分及入选门槛不变。'),('旧版对照','分数定义已改变，旧新分差不能直接解释为技术升降；见名单变化表。')]
 corecols=['report_rank','category_id','name','core_papers','core_active_years','core_active_quarters','core_institutions','recent_papers','share_growth_ratio','core_score','robustness_grade','data_scenarios_passed','time_sensitivity_note','paper_review_reason']
 emcols=['report_rank','category_id','name','recent_papers','baseline_annual_mean_papers','multiyear_share_ratio','share_growth_ratio','historical_peak_ratio','three_year_log_share_slope','emerging_score','robustness_grade','data_scenarios_passed','time_sensitivity_note','paper_review_reason']
 tables=[('核心热点','results/core_hotspots.csv',corecols),('新兴热点','results/emerging_hotspots.csv',emcols),('潜在跟踪方向','results/potential_priority.csv',None),('750类完整结果','results/hotspot_summary_all750.csv',None),('多年方法前后变化','results/multiyear_before_after.csv',None),('五年年度轨迹','results/multiyear_annual_trajectories.csv',None),('核心数值候选待审','results/multiyear_core_candidates.csv',None),('新兴数值候选审阅','results/multiyear_emerging_candidates.csv',None),('跨年补审结论','results/multiyear_semantic_review.csv',None),('跨年已读样本','review/multiyear_displayed_samples.csv',None),('多年数据覆盖','results/multiyear_coverage_audit.csv',None),('数据稳健性分级','results/multiyear_robustness_grades.csv',None),('观察与待复核','results/watch_and_downgraded.csv',None)]
 for name,file in [('潜在统一单元全表','potential_unified_metrics.csv'),('潜在范围定义','potential_unified_scope.csv'),('潜在逐条候选','potential_unified_record_ledger.csv'),('潜在主体证据','potential_unified_applicant_evidence.csv'),('潜在政策任务复核','potential_unified_policy_reviews.csv'),('潜在样本范围复核','potential_unified_document_reviews.csv'),('潜在统计分母','potential_unified_denominators.csv')]:tables.append((name,'results/'+file,None))
 make('750类核心新兴潜在热点分析.xlsx',notes,tables)
 enotes=notes[:1]+[('并列排名','排名时将0—100分取小数点后10位以消除浮点尾差；并列按ID升序展示，实验采用并列最小名次，Spearman也使用舍入后的分数。原始评分及入选门槛不变。'),('总量',f'{len(wt)}次权重扰动，{len(g)}个门槛/数据/窗口场景。'),('消融','移除一个评分成分后归一；资格和语义判断固定。只改分数不改变名单。'),('窗口长度','核心1/3/5年；计数门槛按年缩放，季度按比例缩放。新兴此前2/3/4年基线不与最近一年重叠。'),('截止点','2025Q2、2025Q4、2026Q1；仅对应截止点数值规则，冻结当前标签与语义判断，不输入2026年8月确认；不视为预测回测。'),('删基线年','逐次剔除三个背景年度，重算年均份额、机构与计数检验；最近同比和五年峰值仍固定。'),('数据压力','质量、几何、去重、排除待复核、仅规则支持均重算计数、机构与引用背景；新兴保留主数据方向确认，属于条件测试。'),('排序池','核心/新兴在本次数值资格池内比较，潜在只有4单元；前10/25超过池大小留空。')]
 make('750类热点消融实验与灵敏度分析.xlsx',enotes,[(n,'reliability/'+f,None) for n,f in [('评分指标消融','score_ablation_summary.csv'),('门槛数据窗口灵敏度','gate_and_data_sensitivity.csv'),('1500次权重试验','weight_trials.csv'),('权重排名稳定性','weight_rank_stability.csv'),('逐类别消融排名','score_ablation_ranks.csv'),('情景逐主题依据','data_sensitivity_topic_details.csv'),('潜在统一口径实验','potential_unified_sensitivity_summary.csv'),('潜在逐单元实验依据','potential_unified_sensitivity_details.csv')]])
 dump(BASE/'data/MULTIYEAR_WORKBOOK_TABLES.json',manifest)
 method=read(BASE/'data/MULTIYEAR_METHOD.json');review=pd.read_csv(R/'multiyear_semantic_review.csv');reads=pd.read_csv(BASE/'review/multiyear_displayed_samples.csv')
 r=f'''# 750主题多年核心、新兴与潜在热点初评结果

本版以三年持续表现识别核心，以最近一年相对前三年背景和五年轨迹识别新兴。得到 **{len(core)}个核心、{len(em)}个新兴/持续升温初评候选**。潜在仍是四个任务单元中2项双口径支持、1项条件性跟踪、1项观察。结果供专家评审，不证明技术领先或未来成功。

## 时间窗口及方法

核心主窗2023年7月至2026年6月；最新状态为2025年7月至2026年6月。新兴以最新一年对比2022年7月至2025年6月三个不重叠年度的等权平均份额。五年轨迹覆盖2021年7月至2026年6月。每年分别按当年的合格论文库归一，不将不同覆盖年份直接累加解释为增长。

核心评分：三年平均份额40%、同发表年引用百分位20%、三年去重机构15%、活跃季度占比15%、最近一年份额10%。要求三年≥750篇、每年≥150篇、12季度至少9个季度各≥25篇，最近一年≥250篇且同比份额≥0.8。引用为当前快照累计引用，不是历史引用增长。

新兴评分：对前三年背景的份额增长35%、最近同比增长25%、近三年对数份额趋势20%、同比增长季度占比10%、机构扩展10%。要求最新一年≥100篇、基线年均≥50篇、多年份额比≥1.25、同比≥1.15、至少3个增长季度、三年斜率>0、比前四年最高份额至少高5%。另要求合并基线的计数近似下界>1.05及BH q<0.05；文本日期质量、原几何、全库题名年度去重口径的长期和同比份额，以及最新1至8月同月份额，均≥1.10。计数检验只作筛查，未消除收录偏差、时序相关或标签误差。

排名时将0—100分取小数点后10位以消除浮点尾差；并列按ID升序展示，实验采用并列最小名次，Spearman也使用舍入后的分数。原始评分及入选门槛不变。

主方法和门槛先于查看新名单固定，没有调整参数以保留旧榜。低基数新苗头、未超过历史峰值的恢复、数值通过但语义未审者进入观察/候选表。三年斜率是描述性量；五年峰值判断有门槛敏感性，详见实验。

## 数据覆盖与可比性

{table(cov,['publication_year','papers','observed_months','title_only_fraction','jan1_proxy_fraction','metadata_match_fraction'])}

2026年只有1至8月，上表不能用于全年总量同比。文本完整度和日期集中现象随年份变化，多年份额归一只能控制库规模，不能消除语种、学科或文献类型构成变化。几何及待审标记仍继承原分类，不是750类别准确率。

## 核心候选

{table(core,corecols[:10])}

## 新兴与持续升温候选

{table(em,emcols[:10])}

新兴表示近期文献关注度在多年背景上继续抬升，不要求技术第一次出现。因此成熟工程方向也可能入选，不能把“新兴热点”直接写成“新技术”。

## 样本与新增审阅

本次实际展示阅读{len(reads)}条跨年度题名/最多650字符摘录，涉及{reads.category_id.nunique()}个方向；决定及证据行ID逐项保留。新数值候选先补读三年样本，原核心发现跨年混杂时扩大抽样。未全量人工清洗，不提供类别准确率。语义未通过和未审阅分列，不能把未入选解释为没有研究价值。

{table(review,['category_id','name','core_review_pass','emerging_review_pass','decision','reason'])}

## 与一年主窗版的变化

{table(changes,['family','category_id','name','change'])}

本轮4个原核心退出来自跨年样本发现的范围问题，不是认定这些研究方向不再重要。若只改变数值方法并沿用旧审阅，原15核心和16新兴均仍通过。分数与排名定义发生变化，不能把分数差当作同一量尺下的上升或下降。当前完整数值候选为{int(z.core_eligible.sum())}个核心、{int(z.emerging_robust.sum())}个新兴；主榜只列有样本范围依据且符合对应门槛的方向。核心其他名次候选未全部新增人工审阅。

{main_text(table)}

完整指标、未审候选和证据见[主分析Excel](750类核心新兴潜在热点分析.xlsx)。时间窗、数据和参数依赖见[实验报告](EXPERIMENT_REPORT.md)。
'''
 (BASE/'REPORT.md').write_text(r)
 windows=g[g.kind.isin(['window_length','cutoff_sensitivity','leave_year_out'])];data=g[g.kind.eq('data_sensitivity')]
 exp=f'''# 多年热点分析：消融与灵敏度实验

本版重做核心/新兴评分消融、门槛、时间窗与覆盖压力测试，保留潜在统一任务实验。共{len(wt)}次权重扰动、{len(g)}个情景，其中潜在71个情景。每类500次权重扰动，固定种子75020260925，各原权重乘0.8—1.2后归一。

## 评分与权重

{table(abl[abl.scenario.ne('baseline')],['family','scenario','pool_size','spearman','max_abs_rank_change'])}

排名时将0—100分取小数点后10位以消除浮点尾差；并列按ID升序展示，实验采用并列最小名次，Spearman也使用舍入后的分数。原始评分及入选门槛不变。

排名池为本次数值资格池；得分无硬入选阈值，单独改权重不改变正式资格。500次Spearman最小值：核心{wt[wt.family.eq('core')].spearman.min():.4f}，新兴{wt[wt.family.eq('emerging')].spearman.min():.4f}，潜在{wt[wt.family.eq('potential')].spearman.min():.4f}。稳定性不等于准确率。

## 多时间尺度和历史截止点

{table(windows,['family','scenario','baseline_count','retained','retention','gained','lost'])}

本次核心在1/3/5年口径下均保留12项，新兴在2/3/4年背景下均保留22项；但逐年剔除时蒸汽重整制氢与废塑料回收转化存在未保留情景，已在主表单列提示。较早截止点的新兴入选较少，说明结果具有时期依赖性，不能将本版名单追认成早年已发现的热点。

核心1/3/5年按每年250篇及活跃季度比例缩放，不沿用固定总量门槛。新兴基线2/3/4年均与最近一年不重叠；逐个删去背景年度重算份额、机构和计数检验，最近同比/五年峰值保持。早期截止点只应用对应年份数值条件，不输入后续YTD确认；标签、语义和引用快照仍冻结，因此是回顾性窗口敏感性，不是独立历史预测回测。

![时间长度检验](figures/multiyear_windows.png)

## 数据覆盖与标签压力

{table(data,['family','scenario','baseline_count','retained','retention','lost'])}

每种过滤重算当期背景分母、机构集合与同发表年引用背景；新兴另保持主口径数据方向确认，属于条件压力测试，不声称跑完所有交叉过滤组合。删除待审或仅保留规则支持会改变文献构成，未通过不等于方向无价值。历史更长不自动消除这些偏差。

![数据压力测试](figures/multiyear_data_stress.png)

{experiment_text(table).replace('核心/新兴数值结果保持原口径。','核心/新兴现已改为本报告的多年口径；潜在任务定义及数值保持不变。')}

所有逐次权重、阈值扫描及情景入选依据见[实验Excel](750类热点消融实验与灵敏度分析.xlsx)。本版不以调权稳定代替分类质量验证，也不以多年窗口替代专家评审。
'''
 (BASE/'EXPERIMENT_REPORT.md').write_text(exp)
 (BASE/'README.md').write_text('''# 750主题多年热点分析

- [主分析Excel](750类核心新兴潜在热点分析.xlsx)
- [实验Excel](750类热点消融实验与灵敏度分析.xlsx)
- [报告](REPORT.md) / [实验报告](EXPERIMENT_REPORT.md)

当前方法见data/MULTIYEAR_METHOD.json。核心近三年加最近一年；新兴最近一年对比此前三年非重叠背景并检查五年轨迹。潜在保持四个统一任务单元。

复现顺序：multiyear_prepare.py → multiyear_build.py → multiyear_review_samples.py（展示样本后另写明确判断）→ multiyear_build.py → multiyear_experiments.py → multiyear_finalize.py → multiyear_delivery.py → multiyear_validate.py。复算使用已冻结的语义决定不等于完成新的专家审核。独立仓库提供不依赖全量数据库的汇总回放。

旧一年主窗版在audit/MULTIYEAR_BASELINE.json所指备份中；旧年度脚本不再是当前默认入口。历史截止点为回顾性窗口诊断，不是预测回测。
''')
 save(pd.DataFrame({'field':z.columns,'中文释义':[CN.get(x,x) for x in z.columns]}),'field_dictionary.csv')
 print('MULTIYEAR_DELIVERY',len(core),len(em),len(g),flush=True)
if __name__=='__main__':deliver()
