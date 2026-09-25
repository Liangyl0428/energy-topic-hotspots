"""Independent source counts, input hashes, decisions, experiments and Excel readback."""
from common import *
import pandas as pd,numpy as np,duckdb
from openpyxl import load_workbook
from scipy.stats import spearmanr
from multiyear_scoring import windows,score,gates,ranking_key
from multiyear_delivery import safe,CN
from validate_potential import run as validate_potential

def run():
 checks={};oldroot=Path(read(BASE/'audit/MULTIYEAR_BASELINE.json')['path'])
 for path,h in read(BASE/'audit/INPUT_HASHES.json').items():assert sha(ROOT/path)==h,path
 for i,p in enumerate(read(BASE/'data/ADAPTER_MANIFEST.json')['parts']):
  name=Path(p['file']).name
  for path,key in [(BASE/p['file'],'sha256'),(LABELS/'results/assignments'/name,'source_assignment_sha256'),(CORPUS/'results/merged_assignments'/name,'original_metadata_sha256')]:assert sha(path)==p[key],str(path)
  if i%80==0:print('HASH',i,flush=True)
 checks['all259_source_label_adapter_hashes_unchanged']=True
 z=pd.read_csv(BASE/'results/hotspot_summary_all750.csv').set_index('category_id');assert len(z)==750 and z.index.is_unique
 c=duckdb.connect();c.execute('SET threads=4');c.from_parquet(str(BASE/'data/assignments/*.parquet')).create_view('a');c.execute('ATTACH '+sqlstr(ROOT/'analyze/data/independent_corpus.duckdb')+' AS metadata (READ_ONLY)')
 for i,w in enumerate(windows()):
  start=str(pd.Period(w[0]).start_time.date());end=str(pd.Period(w[-1]).end_time.date())
  v=c.execute("SELECT category_id,count(*) n FROM a WHERE source='paper' AND category_index>=0 AND NOT retracted AND NOT template_record AND date BETWEEN ? AND ? GROUP BY 1",[start,end]).df().set_index('category_id').n.reindex(z.index,fill_value=0)
  np.testing.assert_array_equal(v,z[f'year{i}_papers']);assert v.sum()==z[f'year{i}_denominator'].iloc[0]
  np.testing.assert_allclose((v+.5)/(v.sum()+375),z[f'year{i}_share'],rtol=1e-12)
 checks['five_nonoverlapping_year_blocks_independently_counted']=True
 np.testing.assert_array_equal(z.core_papers,z[['year0_papers','year1_papers','year2_papers']].sum(axis=1))
 np.testing.assert_allclose(z.multiyear_share_ratio,z.year0_share/z[['year1_share','year2_share','year3_share']].mean(axis=1),rtol=1e-12)
 np.testing.assert_allclose(z.historical_peak_ratio,z.year0_share/z[['year1_share','year2_share','year3_share','year4_share']].max(axis=1),rtol=1e-12)
 apps=c.execute("SELECT a.category_id,count(DISTINCT lower(trim(n.name))) AS name_count FROM a JOIN metadata.works w ON substr(a.doc_id,7)=w.work_id CROSS JOIN UNNEST(string_split(coalesce(w.institutions,''),';')) n(name) WHERE a.source='paper' AND a.category_index>=0 AND NOT a.retracted AND NOT a.template_record AND a.date BETWEEN '2023-07-01' AND '2026-06-30' AND trim(n.name) NOT IN ('','None','nan','[]') GROUP BY 1").df().set_index('category_id').name_count.reindex(z.index,fill_value=0)
 np.testing.assert_array_equal(apps,z.core_institutions);checks['three_year_institution_union_recomputed']=True
 for fam,flag in [('core','final_core'),('emerging','final_emerging')]:
  gate=gates(z,fam).all(axis=1)&z[fam+'_review_pass'];assert gate.equals(z[flag])
  df=pd.read_csv(BASE/'results'/f'{fam}_hotspots.csv');assert set(df.category_id)==set(z.index[gate]);checks[fam+'_selected']=len(df)
 # Freeze potential inputs/results, and independently recheck record-level invariants.
 for name in ['potential_unified_metrics.csv','potential_priority.csv','potential_unified_record_ledger.csv','potential_unified_policy_reviews.csv','potential_unified_applicant_evidence.csv']:
  assert sha(BASE/'results'/name)==sha(oldroot/'results'/name),name
 checks['potential']=validate_potential(allow_literature_change=True)
 reviews=pd.read_csv(BASE/'results/multiyear_semantic_review.csv');shown=pd.read_csv(BASE/'review/multiyear_displayed_samples.csv').fillna('');raw=pd.read_parquet(BASE/'evidence/multiyear_review_samples.parquet').set_index('row_id')
 assert len(shown)==192 and shown.row_id.nunique()==192
 for r in shown.itertuples():assert raw.loc[r.row_id,'body'].startswith(r.excerpt) and raw.loc[r.row_id,'category_id']==r.category_id
 for r in reviews.itertuples():assert set(map(int,r.read_row_ids.split(';')))==set(shown[shown.category_id.eq(r.category_id)].row_id)
 checks['cross_year_reviewed_records']=len(shown);checks['cross_year_reviewed_categories']=len(reviews)
 g=pd.read_csv(BASE/'reliability/gate_and_data_sensitivity.csv');wt=pd.read_csv(BASE/'reliability/weight_trials.csv');assert len(wt)==1500 and len(g)==172
 for fam in ['core','emerging']:
  comp=pd.read_csv(BASE/'results'/f'multiyear_{fam}_components.csv').set_index('category_id');mask=z.core_eligible if fam=='core' else z.emerging_robust;idx=z.index[mask]
  for row in wt[wt.family.eq(fam)].itertuples():
   w=pd.Series(json.loads(row.weights));assert np.isclose(w.sum(),1);ss=100*comp.dot(w);rho=spearmanr(ranking_key(z.loc[idx,fam+'_score']),ranking_key(ss.loc[idx])).statistic;assert np.isclose(rho,row.spearman)
  b=g[g.family.eq(fam)&g.scenario.eq('baseline')].iloc[0];assert set(b.selected_ids.split(';'))==set(z.index[z['final_'+fam]])
 checks['all1500_weight_trials_and_baseline_sets']=True;checks['scenarios']=len(g)
 cells=0
 for workbook in read(BASE/'data/MULTIYEAR_WORKBOOK_TABLES.json'):
  wb=load_workbook(BASE/workbook['workbook'],read_only=True,data_only=True)
  for spec in workbook['tables']:
   df=pd.read_csv(BASE/spec['file'])[spec['columns']];ws=wb[spec['sheet']];rows=ws.iter_rows(values_only=True);headers=next(rows);assert list(headers)==[CN.get(x,x) for x in df.columns]
   assert ws.max_row==len(df)+1 and ws.max_column==len(df.columns)
   for actual,expected in zip(rows,df.itertuples(index=False,name=None)):
    for a,e in zip(actual,expected):
     e=safe(e)
     if e is None or e=='':assert a is None or a==''
     elif isinstance(e,(int,float,np.number)) and not isinstance(e,bool):assert a is not None and np.isclose(a,e,rtol=1e-11,atol=1e-11,equal_nan=True)
     else:assert a==e,(spec['sheet'],a,e)
     cells+=1
  wb.close()
 checks['excel_cells_readback']=cells
 checks['passed']=True;checks['completed']=now();dump(BASE/'data/VALIDATION.json',checks);dump(BASE/'data/MULTIYEAR_VALIDATION.json',checks)
 paths=[p for folder in ['results','reliability','src','figures'] for p in (BASE/folder).glob('*') if p.is_file()]+[BASE/f for f in ['REPORT.md','EXPERIMENT_REPORT.md','METHOD.json','750类核心新兴潜在热点分析.xlsx','750类热点消融实验与灵敏度分析.xlsx']]
 dump(BASE/'DELIVERY_MANIFEST.json',dict(version='750-multiyear-v1',created=now(),files=[dict(file=str(p.relative_to(BASE)),bytes=p.stat().st_size,sha256=sha(p)) for p in paths]))
 c.close();print('MULTIYEAR_VALIDATION_PASSED',json.dumps({k:v for k,v in checks.items() if k!='potential'},ensure_ascii=False),flush=True)
 return checks
if __name__=='__main__':run()
