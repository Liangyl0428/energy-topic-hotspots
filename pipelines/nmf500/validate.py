"""Validate generated NMF500 temporal counts, review lineage and cosine assignments."""
from pathlib import Path
import hashlib
import json

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from build import OUTPUT, dump


def validate(out=OUTPUT):
    checks = []
    def check(name, ok):
        checks.append({'check': name, 'passed': bool(ok)})
    papers = pd.read_parquet(out / 'paper_assignments.parquet')
    clean = papers[papers.topic_id.ge(0) & papers.date.le('2026-06-30') & ~papers.is_retracted.fillna(False) & ~papers.template_record.fillna(False)]
    quarters = pd.read_csv(out / 'quarter_counts.csv')
    metrics = pd.read_csv(out / 'hotspot_metrics.csv').set_index('category_id')
    reviewed = pd.read_csv(out / 'reviewed_hotspot_metrics.csv').set_index('category_id')
    potential = pd.read_csv(out / 'reviewed_potential_metrics.csv').set_index('category_id')
    policy = pd.read_csv(out / 'policy_evidence_review.csv').fillna('')
    transfer = pd.read_parquet(out / 'transfer_evidence.parquet').set_index('doc_id')
    assignments = pd.read_parquet(out / 'patent_policy_assignments.parquet')
    summary = json.loads((out / 'SUMMARY.json').read_text())
    check('500_unique_topics', len(metrics) == 500 and metrics.index.is_unique)
    check('paper_ids_unique', papers.work_id.is_unique)
    check('complete_500_by_20_quarters', len(quarters) == 10000 and not quarters.duplicated(['category_id', 'period']).any())
    observed = clean.groupby(['category_id', 'period']).size()
    expected = quarters.set_index(['category_id', 'period']).papers
    check('quarter_counts_reproduce_from_document_rows', expected.equals(observed.reindex(expected.index, fill_value=0).rename('papers')))
    recent = clean[clean.date.between('2025-07-01', '2026-06-30')].groupby('category_id').size().reindex(metrics.index, fill_value=0)
    check('recent_counts_reproduce', np.array_equal(recent, metrics.recent_papers))
    check('sample_totals_match_summary', len(papers) == summary['input_papers'] and len(clean) == summary['eligible_papers'])
    check('rank_scores_finite', np.isfinite(metrics[['core_score', 'emerging_score']]).all().all())
    check('followup_is_subset_of_numeric_candidates', (~reviewed.core_followup | metrics.core_numeric_candidate).all() and (~reviewed.emerging_followup | metrics.emerging_numeric_candidate).all() and (~potential.potential_followup | potential.potential_numeric_candidate).all())
    check('all_selected_new_emerging_preserve_filter_direction', reviewed.loc[reviewed.emerging_followup, 'emerging_robust_candidate'].all())
    check('conditional_core_not_misreported_as_robust', int(reviewed.loc[reviewed.core_followup, 'core_robust_candidate'].sum()) == summary['core_followup_passing_all_data_filters'])
    check('no_unverified_potential_promoted_to_full_verification', not potential.fully_verified_potential_hotspot.any())
    approved_families = policy[policy.accepted.eq(True)].groupby('category_id').family.nunique().reindex(potential.index, fill_value=0)
    check('reviewed_policy_groups_deduplicated', np.array_equal(approved_families, potential.reviewed_policy_family_count))
    excerpts_ok = True
    for row in policy.itertuples():
        body = transfer.loc[row.doc_id, 'body']
        excerpts_ok &= body[row.quote_start:row.quote_end] == row.quote
        excerpts_ok &= hashlib.sha256(row.quote.encode()).hexdigest() == row.quote_sha256
        excerpts_ok &= hashlib.sha256(body.encode()).hexdigest() == row.body_sha256
    check('reviewed_policy_quotes_match_frozen_source_and_hashes', excerpts_ok)
    old = policy[policy.doc_id.eq('policy:bfeb5830e2f4a417e2a38695')]
    check('2006_policy_not_counted_as_recent_support', len(old) == 1 and not old.iloc[0].accepted and old.iloc[0].reviewed_issue_date == '2006-11-03')
    check('transfer_unique_complete_and_finite', assignments.doc_id.is_unique and len(assignments) == summary['input_transfer_documents'] and np.isfinite(assignments[['topic_1_cosine', 'topic_2_cosine', 'topic_3_cosine']]).all().all())
    check('top3_descending_and_valid_ids', (assignments.topic_1_cosine >= assignments.topic_2_cosine).all() and (assignments.topic_2_cosine >= assignments.topic_3_cosine).all() and assignments[['topic_1_id', 'topic_2_id', 'topic_3_id']].ge(0).all().all() and assignments[['topic_1_id', 'topic_2_id', 'topic_3_id']].lt(500).all().all())
    centers = np.load(out / 'topic_centroids.npy')
    norms = np.linalg.norm(centers, axis=1)
    check('centers_finite_unit_norm_or_inactive', centers.shape == (500, 1024) and np.isfinite(centers).all() and np.all(np.isclose(norms, 1, atol=1e-5) | np.isclose(norms, 0)))
    wb = load_workbook(out / '500主题热点审阅结果.xlsx', read_only=True)
    check('review_workbook_matches_followup_counts', wb['核心条件跟踪'].max_row == summary['core_conditional_followup'] + 1 and wb['新兴跟踪'].max_row == summary['emerging_followup'] + 1 and wb['潜在应用线索'].max_row == summary['potential_application_leads'] + 1)
    wb.close()
    result = {'passed': sum(x['passed'] for x in checks), 'total': len(checks), 'checks': checks, 'interpretation': '数值、记录级回算与引文完整性验证；非分类准确率或专家认可率'}
    dump(out / 'VALIDATION.json', result)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if not all(x['passed'] for x in checks):
        raise SystemExit(1)


if __name__ == '__main__':
    validate()
